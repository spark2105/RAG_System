from __future__ import annotations

import atexit
from io import BytesIO
import os
from pathlib import Path
import socket
import subprocess
import sys
import time
from urllib.parse import urlsplit

import httpx
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import streamlit as st

from llm_service import (
    AUTOMATIC_LLM_INPUT_COLUMNS,
    AUTOMATIC_LLM_MODE_FIELDS,
    AUTOMATIC_LLM_OUTPUT_COLUMNS,
    AUTOMATIC_LLM_QUERY_MODES,
    DEFAULT_GEMINI_MODEL,
    GEMINI_MODEL_LABELS,
    GEMINI_MODEL_OPTIONS,
)
from rss_pipeline import AUTOMATIC_RETRIEVAL_MODES, load_pipeline_config


st.set_page_config(page_title="RAG_System", page_icon="🧪", layout="wide")

WORKER_URL = os.getenv("WORKER_URL", "http://localhost:9000").rstrip("/")
PROJECT_ROOT = Path(__file__).resolve().parent


def _api_request(
    method: str,
    path: str,
    *,
    json: dict | None = None,
    timeout: float = 30,
) -> dict:
    try:
        with httpx.Client(base_url=WORKER_URL, timeout=timeout) as client:
            response = client.request(method, path, json=json)
            response.raise_for_status()
            return response.json()
    except httpx.HTTPStatusError as exc:
        try:
            detail = exc.response.json().get("detail", exc.response.text)
        except Exception:
            detail = exc.response.text
        raise RuntimeError(f"Worker request failed ({exc.response.status_code}): {detail}") from exc
    except httpx.HTTPError as exc:
        raise RuntimeError(f"Worker is not reachable at {WORKER_URL}: {exc}") from exc


def _worker_health() -> dict | None:
    try:
        return _api_request("GET", "/health", timeout=5)
    except RuntimeError:
        return None


@st.cache_resource(show_spinner=False)
def _ensure_local_worker() -> subprocess.Popen | None:
    """Start one local worker for direct Streamlit development when needed."""
    worker_address = urlsplit(WORKER_URL)
    worker_host = (worker_address.hostname or "").lower()
    auto_start = os.getenv("AUTO_START_WORKER", "true").strip().lower() not in {"0", "false", "no"}
    if not auto_start or worker_host not in {"localhost", "127.0.0.1"}:
        return None
    worker_port = worker_address.port or 9000
    if _worker_health() is not None:
        return None
    try:
        with socket.create_connection((worker_host, worker_port), timeout=0.5):
            return None
    except OSError:
        pass

    if worker_port != 9000:
        raise RuntimeError("Local worker auto-start expects WORKER_URL to use port 9000.")

    process = subprocess.Popen(
        [
            sys.executable,
            "-m",
            "uvicorn",
            "worker_service:app",
            "--host",
            "127.0.0.1",
            "--port",
            "9000",
        ],
        cwd=PROJECT_ROOT,
        env={**os.environ, "PYTHONUNBUFFERED": "1"},
    )

    def stop_worker() -> None:
        if process.poll() is None:
            process.terminate()

    atexit.register(stop_worker)
    return process


def _pipeline_config() -> dict:
    try:
        return load_pipeline_config()
    except Exception:
        return {
            "embedding_model": "nvidia/Nemotron-3-Embed-1B-BF16",
            "embedding_device": "auto",
            "article_limit_per_feed": 5,
            "article_text_enabled": True,
            "request_timeout_seconds": 20,
            "vector_weight": 0.6,
            "bm25_weight": 0.4,
            "rrf_k": 60,
        }


def _feed_urls() -> list[str]:
    return list(_api_request("GET", "/config/feeds").get("feed_urls", []))


def _save_feed_urls(feed_text: str) -> list[str]:
    urls = []
    seen = set()
    for value in feed_text.splitlines():
        url = value.strip()
        if url and not url.startswith("#") and url not in seen:
            urls.append(url)
            seen.add(url)
    _api_request("PUT", "/config/feeds", json={"feed_urls": urls})
    return urls


def _wait_for_job(job_id: str) -> dict:
    progress = st.progress(0)
    status_box = st.empty()
    started = time.monotonic()
    while True:
        job = _api_request("GET", f"/jobs/{job_id}", timeout=10)
        status = str(job.get("status", "unknown"))
        progress.progress(int(job.get("progress", 0)))
        message = str(job.get("message") or status.title())
        status_box.info(f"Job `{job_id}` · **{status}** · {message}")
        if status in {"succeeded", "failed", "cancelled"}:
            progress.empty()
            status_box.empty()
            if status != "succeeded":
                raise RuntimeError(str(job.get("error") or "Worker job failed."))
            return job
        if time.monotonic() - started > 3600:
            raise RuntimeError("The worker job exceeded the one-hour UI timeout.")
        time.sleep(1)


def _render_health(local_worker: subprocess.Popen | None = None) -> None:
    health = _worker_health()
    if health:
        st.sidebar.success("Worker ready")
    elif local_worker is not None and local_worker.poll() is None:
        st.sidebar.info("Worker starting")
        st.sidebar.caption("The worker is initializing.")
    else:
        st.sidebar.error("Worker unavailable")
        st.sidebar.caption(f"Worker URL: {WORKER_URL}")


def _wait_for_local_worker(local_worker: subprocess.Popen | None) -> None:
    if local_worker is None or local_worker.poll() is not None or _worker_health() is not None:
        return
    st.sidebar.info("Worker starting")
    st.sidebar.caption("The worker is initializing. The application will continue automatically.")
    st.info("The worker is starting. Please wait...")
    time.sleep(1)
    st.rerun()


def _render_ingestion() -> None:
    st.title("RSS Ingestion")
    st.caption("Configure RSS feeds, run ingestion manually, or enable the three-minute scheduler.")

    config = _pipeline_config()
    try:
        existing_urls = _feed_urls()
    except RuntimeError as exc:
        st.error(str(exc))
        existing_urls = []

    feed_text = st.text_area(
        "RSS feed URLs",
        value="\n".join(existing_urls),
        height=150,
        help="Enter one HTTP or HTTPS URL per line.",
    )

    left, middle, right = st.columns(3)
    with left:
        article_limit = st.number_input(
            "Articles per feed",
            min_value=0,
            value=int(config.get("article_limit_per_feed", 5)),
            step=1,
        )
    with middle:
        timeout = st.number_input(
            "Request timeout (seconds)",
            min_value=1,
            value=int(config.get("request_timeout_seconds", 20)),
            step=1,
        )
    with right:
        fetch_full_text = st.checkbox(
            "Fetch full article text",
            value=bool(config.get("article_text_enabled", True)),
        )
    st.caption("Article content is stored after technical HTML extraction only; no semantic rewriting is applied.")

    save_col, manual_col = st.columns(2)
    with save_col:
        if st.button("Save feed configuration", width="stretch"):
            try:
                urls = _save_feed_urls(feed_text)
                st.success(f"Saved {len(urls)} feed(s).")
            except RuntimeError as exc:
                st.error(str(exc))
    with manual_col:
        if st.button("Start manual ingestion", type="primary", width="stretch"):
            try:
                urls = _save_feed_urls(feed_text)
                job = _api_request(
                    "POST",
                    "/jobs/manual-ingestion",
                    json={
                        "feed_urls": urls,
                        "article_limit_per_feed": int(article_limit),
                        "timeout": int(timeout),
                        "fetch_full_text": fetch_full_text,
                    },
                )
                completed = _wait_for_job(str(job["job_id"]))
                result = completed.get("result", {})
                st.success(
                    f"Manual ingestion completed: {result.get('inserted_count', 0)} new, "
                    f"{result.get('skipped_count', 0)} skipped."
                )
                if result.get("errors"):
                    st.warning(f"{len(result['errors'])} feed or article issue(s) were recorded.")
            except RuntimeError as exc:
                st.error(str(exc))

    st.divider()
    st.subheader("Automatic ingestion")
    try:
        scheduler = _api_request("GET", "/automatic-ingestion/status")
        scheduler_enabled = bool(scheduler.get("enabled"))
        status_text = "Running" if scheduler_enabled else "Paused"
        st.info(
            f"Status: **{status_text}** · Interval: {int(scheduler.get('interval_seconds', 180))} seconds"
        )
        if scheduler.get("last_run_at"):
            st.caption(f"Last scheduled run: {scheduler['last_run_at']}")
        auto_start, auto_pause = st.columns(2)
        with auto_start:
            if st.button("Start automatic ingestion", type="primary", disabled=scheduler_enabled):
                _save_feed_urls(feed_text)
                _api_request("POST", "/automatic-ingestion/start")
                st.success("Automatic ingestion started. The first check runs immediately.")
                st.rerun()
        with auto_pause:
            if st.button("Pause automatic ingestion", disabled=not scheduler_enabled):
                _api_request("POST", "/automatic-ingestion/pause")
                st.success("Automatic ingestion paused.")
                st.rerun()
    except RuntimeError as exc:
        st.error(str(exc))


def _render_embeddings() -> None:
    st.title("Embeddings")
    st.caption("Create embeddings manually after RSS ingestion and maintain the vector database.")
    config = _pipeline_config()
    st.code(
        f"model: {config.get('embedding_model', '')}\n"
        f"configured device: {config.get('embedding_device', 'auto')}",
        language="text",
    )

    try:
        health = _worker_health() or {}
        stats = _api_request("GET", "/stats")
        sql_stats = stats.get("sql", {})
        vector_stats = stats.get("vector", {})
        st.info(
            f"Articles: {sql_stats.get('articles', 0)} · Embedded: {sql_stats.get('embedded', 0)} · "
            f"Pending: {sql_stats.get('pending', 0)}"
        )
        st.caption(
            f"Vector rows: {vector_stats.get('rows', 0)} · Dimensions: {vector_stats.get('dimensions', 0)} · "
            f"Full-text index: {'available' if vector_stats.get('fts_index_exists') else 'missing'} · "
            f"Loaded device: {health.get('embedding_device') or 'not loaded'}"
        )
    except RuntimeError as exc:
        st.error(str(exc))

    if st.button("Create embeddings", type="primary"):
        try:
            job = _api_request("POST", "/jobs/embeddings", json={"local_files_only": False})
            completed = _wait_for_job(str(job["job_id"]))
            result = completed.get("result", {})
            st.success(
                f"Embedding job completed: {result.get('stored_count', 0)} vector row(s) stored."
            )
            if result.get("errors"):
                st.warning(f"{len(result['errors'])} article(s) could not be embedded.")
        except RuntimeError as exc:
            st.error(str(exc))

    st.subheader("Clear vector database")
    st.caption("This removes LanceDB vectors and resets SQLite embedding flags. Articles remain stored.")
    confirm_clear = st.checkbox("I understand that vector data will be removed.")
    if st.button("Clear vector database", disabled=not confirm_clear):
        try:
            result = _api_request("POST", "/vector-db/clear")
            st.success(
                f"Removed {result.get('previous_vector_rows', 0)} vector row(s) and reset "
                f"{result.get('reset_sql_rows', 0)} article flag(s)."
            )
        except RuntimeError as exc:
            st.error(str(exc))


def _results_to_xlsx(
    results: list[dict],
    columns: list[str] | tuple[str, ...],
    sheet_title: str = "Retrieval",
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_title
    worksheet.append(list(columns))
    for row in results:
        worksheet.append([row.get(column) for column in columns])
    header_fill = PatternFill("solid", fgColor="315CCE")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for column_index, column_name in enumerate(columns, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(14, min(42, len(column_name) + 4))
    wrap_columns = {
        index
        for index, column_name in enumerate(columns, start=1)
        if column_name in {
            "Comment_Text",
            "Video_Title",
            "Video_Description",
            "ArticleDescription",
            "ArticleText",
            "RAG_Article_Title",
            "RAG_Article_Description",
            "RAG_Article_Text",
            "LLM_Reason",
        }
    }
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_columns)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _dataframe_rows_for_api(dataframe: pd.DataFrame) -> list[dict]:
    """Convert Excel scalar values, including timestamps, to JSON-safe values."""
    rows: list[dict] = []
    for raw_row in dataframe.to_dict(orient="records"):
        row: dict = {}
        for column, value in raw_row.items():
            if value is None:
                row[column] = ""
                continue
            try:
                if bool(pd.isna(value)):
                    row[column] = ""
                    continue
            except (TypeError, ValueError):
                pass
            if isinstance(value, pd.Timestamp):
                row[column] = value.isoformat(sep=" ")
            else:
                row[column] = value
        rows.append(row)
    return rows


def _render_automatic_llm_evaluation() -> None:
    st.title("Automatic LLM Evaluation")
    st.caption("Evaluate every row of an XLSX file with configurable retrieval and Gemini context modes.")

    uploaded_file = st.file_uploader(
        "Input XLSX file",
        type=["xlsx"],
        help="The required columns depend on the selected LLM and retrieval modes.",
    )
    retrieval_mode = st.radio(
        "Retrieval Query Mode",
        options=list(AUTOMATIC_RETRIEVAL_MODES),
        format_func=lambda value: {
            "Q1": "Q1 · Comment text",
            "Q2": "Q2 · Video title + description",
            "Q3": "Q3 · Comment + video fields",
        }[value],
        horizontal=True,
    )
    llm_mode = st.selectbox(
        "LLM Query Mode",
        options=list(AUTOMATIC_LLM_QUERY_MODES),
        format_func=lambda value: {
            "Baseline1": "Baseline1: Comment",
            "Baseline2": "Baseline2: Comment and Platform Context",
            "Rag1": "RAG1: Comment and RAG Article",
            "Rag2": "RAG2: Comment, Platform Context and RAG Article",
            "BestCase1": "BestCase1: Comment, Manually Picked Ideal Article",
            "BestCase2": "BestCase2: Comment, Platform Context, Manually Picked Ideal Article",
            "WorstCase1": "WorstCase1: Comment, Manually Picked Irrelevant Article",
            "WorstCase2": "WorstCase2: Comment, Platform Context, Manually Picked Irrelevant Article",
        }[value],
        help="Select which input fields are sent to Gemini for each row.",
    )
    configured_model = os.getenv("GEMINI_MODEL", DEFAULT_GEMINI_MODEL).strip()
    model_options = list(GEMINI_MODEL_OPTIONS)
    if configured_model and configured_model not in model_options:
        model_options.insert(0, configured_model)
    model = st.selectbox(
        "Gemini model",
        model_options,
        index=model_options.index(configured_model) if configured_model in model_options else 0,
        format_func=lambda value: GEMINI_MODEL_LABELS.get(value, value),
    )

    config = _pipeline_config()
    vector_weight_percent = st.slider(
        "Vector weight",
        min_value=0,
        max_value=100,
        value=int(round(float(config.get("vector_weight", 0.6)) * 100)),
        step=5,
        format="%d%%",
    )
    vector_weight = vector_weight_percent / 100
    bm25_weight = 1 - vector_weight
    st.caption(f"Retrieval weighting: Vector {vector_weight:.0%} · BM25 {bm25_weight:.0%}")
    if llm_mode not in {"Rag1", "Rag2"}:
        st.info("Retrieval is only executed for Rag1 and Rag2. The RAG output columns remain empty for all other modes.")

    required_columns = {"Comment_ID"}
    required_columns.update(
        field for field in AUTOMATIC_LLM_MODE_FIELDS[llm_mode] if not field.startswith("RAG_")
    )
    if llm_mode in {"Rag1", "Rag2"}:
        retrieval_fields = {
            "Q1": {"Comment_Text"},
            "Q2": {"Video_Title", "Video_Description"},
            "Q3": {"Comment_Text", "Video_Title", "Video_Description"},
        }
        required_columns.update(retrieval_fields[retrieval_mode])

    dataframe = None
    missing_columns: list[str] = []
    if uploaded_file is not None:
        try:
            dataframe = pd.read_excel(uploaded_file, engine="openpyxl")
            dataframe.columns = [str(column).strip() for column in dataframe.columns]
            missing_columns = sorted(required_columns - set(dataframe.columns))
        except Exception as exc:
            st.error(f"Could not read XLSX file: {exc}")
        if dataframe is not None and missing_columns:
            st.error("Missing required columns for the selected modes: " + ", ".join(missing_columns))
        elif dataframe is not None:
            st.success(f"Loaded {len(dataframe)} row(s).")
            preview_columns = [column for column in AUTOMATIC_LLM_INPUT_COLUMNS if column in dataframe.columns]
            st.dataframe(dataframe[preview_columns].head(10), hide_index=True, width="stretch")

    api_key = str(st.session_state.get("gemini_api_key") or "").strip()
    if not api_key:
        st.warning("Configure the Gemini API key in the left sidebar before starting the analysis.")
    if not st.button(
        "Start automatic LLM evaluation",
        type="primary",
        disabled=dataframe is None or bool(missing_columns) or not api_key,
    ):
        return

    rows = _dataframe_rows_for_api(dataframe)
    try:
        with st.spinner("Running automatic LLM evaluation ..."):
            job = _api_request(
                "POST",
                "/jobs/llm/automatic-evaluation",
                json={
                    "rows": rows,
                    "input_columns": list(dataframe.columns),
                    "retrieval_mode": retrieval_mode,
                    "llm_query_mode": llm_mode,
                    "vector_weight": vector_weight,
                    "bm25_weight": bm25_weight,
                    "api_key": api_key,
                    "model": model,
                },
                timeout=3600,
            )
            completed = _wait_for_job(str(job["job_id"]))
            result = completed.get("result", {})
    except RuntimeError as exc:
        st.error(str(exc))
        return

    results = result.get("results", [])
    error_count = int(result.get("error_count", 0))
    if error_count:
        st.warning(
            f"Automatic LLM evaluation completed with {error_count} error row(s). "
            "The affected rows are marked as 'Fehler'."
        )
    else:
        st.success(f"Automatic LLM evaluation completed: {len(results)} row(s) processed.")

    st.subheader("Evaluation log")
    evaluation_log = result.get("evaluation_log", [])
    if evaluation_log:
        log_rows = []
        for entry in evaluation_log:
            try:
                duration_ms = float(entry.get("duration_ms", 0))
            except (TypeError, ValueError):
                duration_ms = 0.0
            duration = f"{duration_ms / 1000:.2f} s" if duration_ms >= 1000 else f"{duration_ms:.0f} ms"
            log_rows.append(
                {
                    "Excel-Zeile": entry.get("row_number", "-"),
                    "Comment_ID": entry.get("comment_id", "-"),
                    "Status": entry.get("status", "-"),
                    "Dauer": duration,
                    "Details": entry.get("detail", ""),
                }
            )
        st.dataframe(log_rows, hide_index=True, width="stretch")
    else:
        st.info("Für diese Evaluation wurden keine Logeinträge zurückgegeben.")

    st.subheader("Results")
    output_columns = result.get("output_columns") or list(AUTOMATIC_LLM_OUTPUT_COLUMNS)
    st.dataframe(
        pd.DataFrame(results, columns=output_columns),
        hide_index=True,
        width="stretch",
    )
    st.download_button(
        "Download XLSX results",
        data=_results_to_xlsx(
            results,
            columns=output_columns,
            sheet_title="LLM Evaluation",
        ),
        file_name=f"LLM_Evaluation_{llm_mode}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )


def _render_manual_llm_evaluation() -> None:
    st.title("Manual LLM Evaluation")
    st.caption("Evaluate one comment with the Gemini API.")

    api_key = str(st.session_state.get("gemini_api_key") or "").strip()
    if not api_key:
        st.warning("Configure the Gemini API key in the left sidebar before starting the evaluation.")
    configured_model = os.getenv("GEMINI_MODEL", DEFAULT_GEMINI_MODEL).strip()
    model_options = list(GEMINI_MODEL_OPTIONS)
    if configured_model and configured_model not in model_options:
        model_options.insert(0, configured_model)
    model = st.selectbox(
        "Gemini model",
        model_options,
        index=model_options.index(configured_model) if configured_model in model_options else 0,
        format_func=lambda value: GEMINI_MODEL_LABELS.get(value, value),
        help="The selected model is used for the next evaluation request.",
    )
    comment = st.text_area("Comment", height=140)
    left, right = st.columns(2)
    with left:
        video_title = st.text_input("Video Title")
        video_description = st.text_area("Video Description", height=120)
        video_upload_time = st.text_input("Video Upload Time")
    with right:
        article_title = st.text_input("Article Title")
        article_description = st.text_area("Article Description", height=120)
        article_text = st.text_area("Article Text", height=260)

    payload = {
        "comment": comment,
        "video_title": video_title,
        "video_description": video_description,
        "video_upload_time": video_upload_time,
        "article_title": article_title,
        "article_description": article_description,
        "article_text": article_text,
    }

    if st.button("Gesamtprompt anzeigen", disabled=not comment.strip()):
        try:
            prompt_result = _api_request(
                "POST",
                "/llm/manual-evaluation/prompt",
                json=payload,
                timeout=30,
            )
            st.session_state["manual_llm_prompt"] = prompt_result["prompt"]
            st.session_state["manual_llm_prompt_payload"] = payload
            st.session_state["manual_llm_prompt_confirmed"] = False
        except (KeyError, RuntimeError) as exc:
            st.error(f"Gesamtprompt konnte nicht geladen werden: {exc}")

    prompt = st.session_state.get("manual_llm_prompt")
    prompt_payload = st.session_state.get("manual_llm_prompt_payload")
    prompt_matches = prompt_payload == payload
    if not prompt:
        st.info("Fülle mindestens das Pflichtfeld Comment aus und zeige anschließend den Gesamtprompt an.")
        return

    st.subheader("Gesamtprompt")
    st.code(prompt, language="text")
    if not prompt_matches:
        st.warning("Die Eingabefelder wurden verändert. Zeige den Gesamtprompt erneut an, bevor du die Evaluation startest.")

    confirmed = st.checkbox(
        "Ich habe den Gesamtprompt geprüft und bestätige die Evaluation.",
        key="manual_llm_prompt_confirmed",
        disabled=not prompt_matches,
    )
    if not st.button(
        "Evaluation starten",
        type="primary",
        disabled=not confirmed or not prompt_matches or not api_key,
    ):
        return
    try:
        with st.spinner("Evaluating comment with Gemini ..."):
            result = _api_request(
                "POST",
                "/llm/manual-evaluation",
                json={**payload, "api_key": api_key, "model": model},
                timeout=360,
            )
    except RuntimeError as exc:
        st.error(str(exc))
        return

    st.subheader("Evaluationslog")
    evaluation_log = result.get("evaluation_log", [])
    if evaluation_log:
        log_rows = []
        for entry in evaluation_log:
            try:
                duration_ms = float(entry.get("duration_ms", 0))
            except (TypeError, ValueError):
                duration_ms = 0.0
            duration = f"{duration_ms / 1000:.2f} s" if duration_ms >= 1000 else f"{duration_ms:.0f} ms"
            log_rows.append(
                {
                    "Schritt": entry.get("step", "-"),
                    "Status": entry.get("status", "-"),
                    "Dauer": duration,
                    "Details": entry.get("detail", ""),
                }
            )
        st.dataframe(log_rows, hide_index=True, width="stretch")
    else:
        st.info("Für diese Evaluation wurden keine Logeinträge zurückgegeben.")

    st.subheader("LLM-Output")
    st.caption(f"Gemini model: {result.get('model', model)}")
    st.code(str(result.get("llm_output") or ""), language="text")
    st.subheader("Bewertung")
    score = result.get("result")
    if isinstance(score, int) and 0 <= score <= 10:
        st.metric("Hassrede-Score", f"{score}/10")
    else:
        st.error("The worker returned an invalid hate speech score.")


def _render_articles() -> None:
    st.title("Article Inventory")
    st.caption("Inspect stored articles and embedding state.")
    try:
        rows = _api_request("GET", "/articles?limit=200").get("articles", [])
    except RuntimeError as exc:
        st.error(str(exc))
        return
    if not rows:
        st.info("No articles are stored yet. Start RSS ingestion first.")
        return
    display_rows = [
        {
            "Title": row.get("Article_Title") or "",
            "Feed": row.get("Feed_URL") or "",
            "Published": row.get("published_date") or "",
            "Embedded": bool(row.get("embedded")),
            "URL": row.get("Article_URL") or "",
        }
        for row in rows
    ]
    st.dataframe(display_rows, hide_index=True, width="stretch", height=430)
    selected_index = st.selectbox(
        "Select an article",
        list(range(len(rows))),
        format_func=lambda index: f"{index + 1}. {rows[index].get('Article_Title') or rows[index].get('Article_URL')}",
    )
    selected = rows[selected_index]
    with st.expander("Selected article", expanded=True):
        st.write(f"**Title:** {selected.get('Article_Title') or '-'}")
        st.write(f"**URL:** {selected.get('Article_URL') or '-'}")
        st.write(f"**Description:** {selected.get('Article_Description') or '-'}")
        st.text_area("Article text", value=selected.get("Article_Text") or "", height=260)


def _render_jobs() -> None:
    st.title("Jobs")
    st.caption("Inspect completed, failed, and active worker jobs. API keys are not stored in job records.")

    if st.button("Refresh jobs"):
        st.rerun()
    try:
        jobs = _api_request("GET", "/jobs?limit=500").get("jobs", [])
    except RuntimeError as exc:
        st.error(str(exc))
        return

    if not jobs:
        st.info("No jobs found.")
        return

    display_rows = [
        {
            "Job ID": job.get("job_id", ""),
            "Type": job.get("job_type", ""),
            "Status": job.get("status", ""),
            "Progress": f"{int(job.get('progress', 0))}%",
            "Message": job.get("message", "") or "",
            "Created": job.get("created_at", "") or "",
            "Started": job.get("started_at", "") or "",
            "Finished": job.get("finished_at", "") or "",
            "Error": job.get("error", "") or "",
        }
        for job in jobs
    ]
    st.dataframe(display_rows, hide_index=True, width="stretch")

    job_ids = [str(job.get("job_id")) for job in jobs]
    selected_job_id = st.selectbox("Job to delete", job_ids)
    selected_job = next((job for job in jobs if str(job.get("job_id")) == selected_job_id), {})
    selected_status = str(selected_job.get("status", ""))
    confirm_selected = st.checkbox(
        "I understand that the selected job record and its stored result will be deleted.",
        key="confirm_delete_selected_job",
    )
    if st.button(
        "Delete selected job",
        disabled=not confirm_selected or selected_status in {"pending", "running"},
    ):
        try:
            _api_request("DELETE", f"/jobs/{selected_job_id}")
            st.success("Job deleted.")
            st.rerun()
        except RuntimeError as exc:
            st.error(str(exc))

    st.divider()
    confirm_all = st.checkbox(
        "I understand that all finished job records and stored results will be deleted.",
        key="confirm_delete_finished_jobs",
    )
    if st.button("Delete all finished jobs", disabled=not confirm_all):
        try:
            result = _api_request("DELETE", "/jobs")
            st.success(f"Deleted {result.get('deleted_count', 0)} finished job(s).")
            st.rerun()
        except RuntimeError as exc:
            st.error(str(exc))


def main() -> None:
    local_worker = _ensure_local_worker()
    _wait_for_local_worker(local_worker)
    _render_health(local_worker)
    with st.sidebar:
        st.header("Navigation")
        page = st.radio(
            "Section",
            [
                "RSS Ingestion",
                "Embeddings",
                "Article Inventory",
                "Jobs",
                "Automatic LLM Evaluation",
                "Manual LLM Evaluation",
            ],
        )
        st.divider()
        st.subheader("Gemini configuration")
        st.text_input(
            "Gemini API key",
            type="password",
            key="gemini_api_key",
            help="The key is kept in this Streamlit session and sent only for evaluation requests.",
        )
        st.caption(
            "API key status: configured"
            if str(st.session_state.get("gemini_api_key") or "").strip()
            else "API key status: not configured"
        )
        st.divider()
        st.caption("RSS ingestion, embeddings, jobs, and LLM-based evaluation.")

    if page == "RSS Ingestion":
        _render_ingestion()
    elif page == "Embeddings":
        _render_embeddings()
    elif page == "Article Inventory":
        _render_articles()
    elif page == "Jobs":
        _render_jobs()
    elif page == "Automatic LLM Evaluation":
        _render_automatic_llm_evaluation()
    else:
        _render_manual_llm_evaluation()


if __name__ == "__main__":
    main()
